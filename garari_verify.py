"""Static LISP and equipment master workbook integrity check."""
from pathlib import Path
from datetime import date, datetime
from openpyxl import load_workbook
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = Path(__file__).resolve().parent
for name in ('08_garari_table.lsp',):
    text = (ROOT / name).read_text(encoding='utf-8-sig')
    depth = 0
    quoted = escaped = comment = False
    for ch in text:
        if comment:
            if ch == '\n':
                comment = False
            continue
        if quoted:
            if escaped:
                escaped = False
            elif ch == '\\':
                escaped = True
            elif ch == '"':
                quoted = False
            continue
        if ch == ';':
            comment = True
        elif ch == '"':
            quoted = True
        elif ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
            assert depth >= 0, f'{name}: extra close parenthesis'
    assert depth == 0 and not quoted, f'{name}: unbalanced form ({depth})'
    print(f'{name}: balanced forms')
MASTER_NAME = '機器マスター_空調.xlsx'
book = load_workbook(ROOT / MASTER_NAME, data_only=False)
assert {'KUCHO_KIKI_MASTER', 'KANKI_MASTER', 'GARARI_MASTER', 'ROOM_MASTER'} <= set(book.sheetnames)
def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
for name, keys in (
    ('GARARI_MASTER', ['SHAPE', 'DIA', 'FD', 'MATERIAL', 'MESH', 'MESH_SPEC', 'FD_TEMP', 'SOUND']),
    ('ROOM_MASTER', ['MODEL']),
):
    sheet = book[name]
    headers = [cell_text(c.value).upper() for c in sheet[1]]
    populated = [h for h in headers if h]
    assert len(populated) == len(set(populated)), f'{name}: duplicate header'
    required = set(keys) | {'MODEL', 'MANUFACTURER', 'DIA', 'SPEC_NOTE', 'APPROVED', 'SOURCE_URL', 'CHECKED_DATE'}
    assert required <= set(headers), f'{name}: missing required headers'
    rows = [dict(zip(headers, map(cell_text, values))) for values in sheet.iter_rows(min_row=3, values_only=True)]
    rows = [row for row in rows if row['MODEL']]
    assert rows, f'{name}: no equipment rows'
    seen = set()
    for row in rows:
        key = tuple(row[k] for k in keys)
        assert key not in seen, f'{name}: duplicate key {key}'
        seen.add(key)
        assert row['APPROVED'] in ('0', '1')
        assert row['MODEL'] and row['DIA'].isdigit() and int(row['DIA']) > 0
        assert not any('\n' in v or '\r' in v for v in row.values())
        if name == 'GARARI_MASTER':
            assert row['SHAPE'] in ('深型', '平型')
            assert all(row[key] in ('あり', 'なし') for key in ('FD', 'MESH', 'SOUND'))
            assert row['MATERIAL']
            assert (row['FD_TEMP'].isdigit() if row['FD'] == 'あり' else row['FD_TEMP'] == '')
            assert (row['MESH_SPEC'].isdigit() if row['MESH'] == 'あり' else row['MESH_SPEC'] == '')
    print(f'{name}: {len(rows)} unique valid workbook rows')
book.close()
